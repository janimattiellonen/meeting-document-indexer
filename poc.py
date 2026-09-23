"""Proof of concept: extract structured meeting data from a directory of documents.

Usage:
    uv run poc.py <directory> [--model qwen3.8:27b-mlx] [--json]

Everything runs locally: text is extracted with PyMuPDF / python-docx / textutil
and sent only to the local Ollama server (http://localhost:11434).
"""

import argparse
import json
import subprocess
import sys
import time
from pathlib import Path

import docx
import pymupdf
import ollama
from pydantic import BaseModel, Field

SUPPORTED = {".pdf", ".docx", ".doc", ".rtf", ".odt"}
MIN_TEXT_CHARS = 50  # below this we assume a scanned PDF without a text layer


class Person(BaseModel):
    name: str
    role: str | None = Field(None, description="esim. puheenjohtaja, sihteeri")


class Topic(BaseModel):
    title: str
    summary: str = Field(description="1-2 lauseen kuvaus käsittelystä ja mahdollisista päätöksistä")


class Meeting(BaseModel):
    meeting_title: str = Field(description="esim. 'Hallituksen kokous 4/2026' tai 'Syyskokous 2025'")
    date: str | None = Field(None, description="Kokouksen päivämäärä muodossa YYYY-MM-DD")
    time: str | None = Field(None, description="Kellonaika, esim. 18.05")
    location: str | None
    present: list[Person]
    absent: list[Person]
    topics: list[Topic]
    summary: str = Field(description="3-5 lauseen yhteenveto koko kokouksesta")


SYSTEM_PROMPT = """Olet avustaja, joka poimii tietoja suomenkielisistä yhdistyksen kokouspöytäkirjoista.
Palauta vain dokumentissa oleva tieto. Älä keksi mitään. Jos tietoa ei ole, käytä null tai tyhjää listaa.
Päivämäärät ovat suomalaisessa muodossa (pp.kk.vvvv); muunna ne muotoon YYYY-MM-DD.
Ota mukaan kaikki käsitellyt asiakohdat, myös muodolliset (avaus, esityslista jne.).
Kirjoita tiivistelmät suomeksi."""


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with pymupdf.open(path) as pdf:
            return "\n".join(page.get_text() for page in pdf)
    if suffix == ".docx":
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    # .doc / .rtf / .odt: macOS built-in converter
    result = subprocess.run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        capture_output=True, text=True, check=True,
    )
    return result.stdout


def analyze(text: str, model: str) -> Meeting:
    response = ollama.chat(
        model=model,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": f"Pöytäkirja:\n\n{text}"},
        ],
        format=Meeting.model_json_schema(),
        think=False,
        options={"temperature": 0, "num_ctx": 32768},
    )
    return Meeting.model_validate_json(response.message.content)


def people(persons: list[Person]) -> str:
    return ", ".join(f"{p.name} ({p.role})" if p.role else p.name for p in persons) or "-"


def print_meeting(path: Path, m: Meeting) -> None:
    print("=" * 80)
    print(f"{m.meeting_title}")
    print(f"Tiedosto:  {path}")
    print(f"Aika:      {m.date or '?'} {m.time or ''}".rstrip())
    print(f"Paikka:    {m.location or '-'}")
    print(f"Läsnä:     {people(m.present)}")
    print(f"Poissa:    {people(m.absent)}")
    print("Aiheet:")
    for i, t in enumerate(m.topics, 1):
        print(f"  {i:>2}. {t.title}: {t.summary}")
    print(f"Yhteenveto: {m.summary}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("directory", type=Path)
    parser.add_argument("--model", default="qwen3.8:27b-mlx")
    parser.add_argument("--json", action="store_true", help="print results as JSON instead of text")
    args = parser.parse_args()

    files = sorted(p for p in args.directory.rglob("*") if p.suffix.lower() in SUPPORTED and not p.name.startswith("~$"))
    if not files:
        sys.exit(f"No supported documents found in {args.directory}")

    results = []
    for path in files:
        print(f"Processing {path.name} ...", file=sys.stderr)
        try:
            text = extract_text(path)
        except Exception as e:
            print(f"  ! could not read: {e}", file=sys.stderr)
            continue
        if len(text.strip()) < MIN_TEXT_CHARS:
            print("  ! no text layer (scanned document?), skipping - needs OCR", file=sys.stderr)
            continue
        started = time.monotonic()
        try:
            meeting = analyze(text, args.model)
        except Exception as e:
            print(f"  ! model failed: {e}", file=sys.stderr)
            continue
        print(f"  done in {time.monotonic() - started:.1f}s", file=sys.stderr)
        if args.json:
            results.append({"file": str(path.resolve()), **meeting.model_dump()})
        else:
            print_meeting(path.resolve(), meeting)

    if args.json:
        print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
