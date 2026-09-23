"""Text extraction. Documents are generated in the test, with invented content."""

import shutil
import subprocess
from pathlib import Path

import docx
import pymupdf
import pytest

from meeting_indexer.extract import discover, extract_pages, has_text, sha256

TEXT = "Hallituksen kokous 1/2007\nLäsnä: Maija Meikäläinen (pj), Teppo Testaaja"


def make_pdf(path: Path, pages: list[str]) -> Path:
    pdf = pymupdf.open()
    for text in pages:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text, fontname="helv")
    pdf.save(path)
    return path


def test_discover_finds_supported_files_recursively(tmp_path: Path) -> None:
    for name in ["2019/a.pdf", "2019/b.DOC", "c.docx", "notes.txt", "~$lock.docx", ".hidden.pdf"]:
        (tmp_path / name).parent.mkdir(parents=True, exist_ok=True)
        (tmp_path / name).touch()

    found = [p.relative_to(tmp_path).as_posix() for p in discover(tmp_path)]

    assert found == ["2019/a.pdf", "2019/b.DOC", "c.docx"]


def test_pdf_text_is_extracted_per_page(tmp_path: Path) -> None:
    path = make_pdf(tmp_path / "a.pdf", ["Sivu yksi", "Sivu kaksi"])
    pages = extract_pages(path)
    assert [p.strip() for p in pages] == ["Sivu yksi", "Sivu kaksi"]


def test_docx_paragraphs_and_tables_are_extracted(tmp_path: Path) -> None:
    document = docx.Document()
    for line in TEXT.splitlines():
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Päätös"
    table.rows[0].cells[1].text = "Hyväksyttiin"
    path = tmp_path / "a.docx"
    document.save(str(path))

    [text] = extract_pages(path)

    assert "Maija Meikäläinen" in text
    assert "Päätös | Hyväksyttiin" in text


@pytest.mark.skipif(shutil.which("textutil") is None, reason="macOS textutil not available")
def test_legacy_doc_is_extracted_with_finnish_characters(tmp_path: Path) -> None:
    html = tmp_path / "a.html"
    body = "".join(f"<p>{line}</p>" for line in TEXT.splitlines())
    html.write_text(f'<html><head><meta charset="utf-8"></head><body>{body}</body></html>', encoding="utf-8")
    doc = tmp_path / "a.doc"
    subprocess.run(["textutil", "-convert", "doc", str(html), "-output", str(doc)], check=True)

    [text] = extract_pages(doc)

    assert "Läsnä: Maija Meikäläinen (pj), Teppo Testaaja" in text


def test_scanned_document_has_no_text(tmp_path: Path) -> None:
    assert not has_text(extract_pages(make_pdf(tmp_path / "scan.pdf", ["", ""])))
    assert has_text([TEXT])


def test_sha256_changes_with_content(tmp_path: Path) -> None:
    a, b = tmp_path / "a", tmp_path / "b"
    a.write_bytes(b"one")
    b.write_bytes(b"two")
    assert sha256(a) != sha256(b)
    assert sha256(a) == "7692c3ad3540bb803c020b3aee66cd8887123234ea0c6e7143c0add73ff431ed"
