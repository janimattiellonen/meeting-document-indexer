"""Finding documents and extracting their text."""

import hashlib
import shutil
import subprocess
from collections.abc import Iterator
from pathlib import Path

import docx
import pymupdf

SUPPORTED_SUFFIXES = {".pdf", ".doc", ".docx", ".rtf", ".odt"}
MIN_TEXT_CHARS = 50  # below this a document is treated as having no text layer (a scan)

SOFFICE_CANDIDATES = ["soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"]


class ExtractionError(Exception):
    pass


def discover(root: Path) -> Iterator[Path]:
    """Yield supported documents under root, sorted, skipping Office lock files and hidden files."""
    for path in sorted(root.rglob("*")):
        if (
            path.is_file()
            and path.suffix.lower() in SUPPORTED_SUFFIXES
            and not path.name.startswith(("~$", "."))
        ):
            yield path


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 16), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_pages(path: Path) -> list[str]:
    """Text of the document, one string per page. Word formats have no pages and return one."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with pymupdf.open(path) as pdf:
            return [str(page.get_text()) for page in pdf]
    if suffix == ".docx":
        return [_docx_text(path)]
    return [_office_text(path)]


def has_text(pages: list[str]) -> bool:
    return sum(len(p.strip()) for p in pages) >= MIN_TEXT_CHARS


def _docx_text(path: Path) -> str:
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _office_text(path: Path) -> str:
    """.doc / .rtf / .odt: macOS textutil first, LibreOffice as a fallback for files it can't read."""
    try:
        result = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            capture_output=True,
            text=True,
            check=True,
            timeout=60,
        )
        if result.stdout.strip():
            return result.stdout
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
        pass

    soffice = next((c for c in SOFFICE_CANDIDATES if shutil.which(c)), None)
    if soffice is None:
        raise ExtractionError("textutil could not read the file and LibreOffice is not installed")
    try:
        result = subprocess.run(
            [soffice, "--headless", "--cat", str(path)],
            capture_output=True,
            text=True,
            check=True,
            timeout=120,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        raise ExtractionError(f"neither textutil nor LibreOffice could read the file: {e}") from e
    return result.stdout.lstrip("﻿")
